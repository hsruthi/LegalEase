from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

from .common import ensure_output_directory, create_filename, clean_text


def generate_docx(title, content, output_directory="generated_documents"):
    """Generate a DOCX document from the supplied title and content."""

    output_dir = ensure_output_directory(output_directory)
    filename = create_filename("LegalEase_Document", "docx")
    file_path = Path(output_dir) / filename

    document = Document()

    # Default font
    styles = document.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(11)

    # Title
    title_paragraph = document.add_paragraph()
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    title_run = title_paragraph.add_run(clean_text(title))
    title_run.bold = True
    title_run.font.size = Pt(16)

    # Content
    if isinstance(content, list):
        for item in content:
            if isinstance(item, dict):
                heading = clean_text(item.get("heading", ""))
                text = clean_text(item.get("text", ""))

                if heading:
                    heading_paragraph = document.add_paragraph()
                    heading_run = heading_paragraph.add_run(heading)
                    heading_run.bold = True
                    heading_run.font.size = Pt(13)

                if text:
                    document.add_paragraph(text)

            else:
                document.add_paragraph(clean_text(item))

    else:
        paragraphs = str(content).split("\n")

        for paragraph in paragraphs:
            paragraph = clean_text(paragraph)

            if paragraph:
                document.add_paragraph(paragraph)

    document.save(file_path)

    return file_path